#!/usr/bin/env python3
"""Google NotebookLM 研修スライド（10枚）PPTX生成スクリプト"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Google brand colors
BLUE = RGBColor(0x42, 0x85, 0xF4)
GREEN = RGBColor(0x0F, 0x9D, 0x58)
YELLOW = RGBColor(0xF4, 0xB4, 0x00)
RED = RGBColor(0xDB, 0x44, 0x37)
TEAL = RGBColor(0x00, 0x89, 0x7B)
PURPLE = RGBColor(0x7B, 0x1F, 0xA2)
INDIGO = RGBColor(0x3F, 0x51, 0xB5)
DARK = RGBColor(0x26, 0x32, 0x38)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT_GRAY = RGBColor(0x99, 0x99, 0x99)
BG_LIGHT = RGBColor(0xF8, 0xF9, 0xFA)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def add_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_gradient_rect(slide, left, top, width, height, color1, color2):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.line.fill.background()
    fill = shape.fill
    fill.gradient()
    fill.gradient_stops[0].color.rgb = color1
    fill.gradient_stops[1].color.rgb = color2


def add_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_rounded_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_circle(slide, left, top, size, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_textbox(slide, left, top, width, height, text, font_size=18,
                color=DARK, bold=False, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_slide_number(slide, num):
    add_textbox(slide, Inches(12.5), Inches(7.0), Inches(0.8), Inches(0.4),
                f"{num} / 10", font_size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.RIGHT)


def add_header_bar(slide, title, color):
    bar = add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(1.1), color)
    add_textbox(slide, Inches(0.6), Inches(0.15), Inches(12), Inches(0.8),
                title, font_size=32, color=WHITE, bold=True)


def add_multi_text(slide, left, top, width, height, lines, font_size=16, color=DARK,
                   line_spacing=1.5, bold=False):
    """Add a textbox with multiple paragraphs."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.bold = bold
        p.font.name = "Calibri"
        p.space_after = Pt(font_size * (line_spacing - 1))
    return txBox


# ============================================================
# SLIDE 1: タイトル
# ============================================================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
add_gradient_rect(slide1, 0, 0, SLIDE_W, SLIDE_H, BLUE, GREEN)

add_textbox(slide1, Inches(1), Inches(1.2), Inches(11.3), Inches(1.0),
            "📓", font_size=60, color=WHITE, alignment=PP_ALIGN.CENTER)
add_textbox(slide1, Inches(1), Inches(2.3), Inches(11.3), Inches(1.2),
            "Google NotebookLM 研修", font_size=48, color=WHITE, bold=True,
            alignment=PP_ALIGN.CENTER)
add_textbox(slide1, Inches(1), Inches(3.6), Inches(11.3), Inches(0.8),
            "〜 生成AIで変わる情報整理・ナレッジ活用の新しいカタチ 〜",
            font_size=22, color=WHITE, alignment=PP_ALIGN.CENTER)
add_textbox(slide1, Inches(1), Inches(5.5), Inches(11.3), Inches(0.5),
            "2026年3月  ｜  社内研修資料",
            font_size=16, color=RGBColor(0xDD, 0xDD, 0xDD), alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 2: アジェンダ
# ============================================================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide2, WHITE)
# Side bar
add_gradient_rect(slide2, 0, 0, Inches(0.12), SLIDE_H, BLUE, GREEN)

add_textbox(slide2, Inches(0.8), Inches(0.5), Inches(5), Inches(0.8),
            "本日のアジェンダ", font_size=36, color=BLUE, bold=True)

agenda_items = [
    "NotebookLM とは？",
    "主要な機能紹介",
    "基本的な使い方（4ステップ）",
    "ビジネス活用事例",
    "NotebookLM の強み",
    "料金プラン",
    "利用時の注意点",
    "まとめ & ネクストステップ",
]
for i, item in enumerate(agenda_items):
    y = Inches(1.6 + i * 0.7)
    circle = add_circle(slide2, Inches(0.9), y, Inches(0.45), BLUE)
    # Number in circle
    tf = circle.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = str(i + 1)
    p.font.size = Pt(16)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].space_before = Pt(0)

    add_textbox(slide2, Inches(1.6), y + Emu(Inches(0.05).emu), Inches(8), Inches(0.4),
                item, font_size=20, color=DARK)
    # Divider line
    if i < len(agenda_items) - 1:
        line = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                       Inches(0.9), y + Inches(0.55), Inches(9), Pt(1))
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(0xEE, 0xEE, 0xEE)
        line.line.fill.background()

add_slide_number(slide2, 2)

# ============================================================
# SLIDE 3: NotebookLMとは
# ============================================================
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide3, WHITE)
add_header_bar(slide3, "NotebookLM とは？", BLUE)

# Left column - description
add_textbox(slide3, Inches(0.6), Inches(1.4), Inches(6), Inches(0.6),
            "Googleが提供するAIナレッジアシスタント", font_size=24, color=BLUE, bold=True)

desc_lines = [
    "・ユーザーがアップロードした資料だけを情報源として、",
    "  AIが要約・質問応答・分析を行うツール",
    "",
    "・最新のGeminiモデルを搭載",
    "・100万トークンのコンテキストウィンドウ",
    "",
    "・Googleアカウントがあれば無料で利用開始可能",
    "  → 導入ハードルが非常に低い",
]
add_multi_text(slide3, Inches(0.6), Inches(2.2), Inches(6), Inches(4.5),
               desc_lines, font_size=16, color=GRAY, line_spacing=1.3)

# Right column - supported sources
sources = [
    ("📄", "PDF / Word / テキスト"),
    ("🌐", "Webサイト / URL"),
    ("🎥", "YouTube動画"),
    ("🎵", "音声ファイル（MP3等）"),
    ("📊", "Google Workspace連携"),
    ("🖼️", "画像ファイル（JPEG/PNG）"),
]
for i, (icon, label) in enumerate(sources):
    y = Inches(1.5 + i * 0.85)
    card = add_rounded_rect(slide3, Inches(7.2), y, Inches(5.5), Inches(0.7),
                            RGBColor(0xF0, 0xF7, 0xFF))
    # Blue left border
    add_rect(slide3, Inches(7.2), y, Inches(0.06), Inches(0.7), BLUE)
    add_textbox(slide3, Inches(7.4), y + Inches(0.1), Inches(0.6), Inches(0.5),
                icon, font_size=22, alignment=PP_ALIGN.CENTER)
    add_textbox(slide3, Inches(8.1), y + Inches(0.12), Inches(4.4), Inches(0.5),
                label, font_size=17, color=DARK, bold=True)

add_slide_number(slide3, 3)

# ============================================================
# SLIDE 4: 主要機能
# ============================================================
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide4, WHITE)
add_header_bar(slide4, "主要な機能紹介", GREEN)

features = [
    ("🎙️", "音声解説", "(Audio Overview)", "AIホスト2人が資料を対話形式で\n解説する音声を自動生成。\n80言語以上対応。", BG_LIGHT, None),
    ("🎬", "動画解説", "(Video Overview)", "資料の要点をビジュアル付き\n解説動画として自動生成。\nYouTubeやSNS共有にも便利。", BG_LIGHT, None),
    ("🧠", "マインドマップ", "", "資料の内容を視覚的に整理。\n複雑な情報の関係性を\n一目で把握。", RGBColor(0xE8, 0xF5, 0xE9), GREEN),
    ("📊", "インフォグラフィック", "生成", "ソース資料を「一目でわかる」\n図解に自動変換。\n2025年12月追加の新機能。", BG_LIGHT, None),
    ("📑", "スライド生成", "& エクスポート", "プロンプト指示で個別スライド\nまで微修正可能。PowerPoint\n形式でダウンロード可。", BG_LIGHT, None),
    ("🔍", "Deep Research", "", "複雑なテーマについて多角的な\n調査を自動実行。統合レポート\nを数分で生成。", BG_LIGHT, None),
]

for i, (icon, title, subtitle, desc, bg_color, border_color) in enumerate(features):
    col = i % 3
    row = i // 3
    x = Inches(0.5 + col * 4.2)
    y = Inches(1.4 + row * 3.0)

    card = add_rounded_rect(slide4, x, y, Inches(3.8), Inches(2.7), bg_color)
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(2)

    add_textbox(slide4, x, y + Inches(0.2), Inches(3.8), Inches(0.5),
                icon, font_size=36, alignment=PP_ALIGN.CENTER)
    add_textbox(slide4, x, y + Inches(0.75), Inches(3.8), Inches(0.4),
                title, font_size=18, color=DARK, bold=True, alignment=PP_ALIGN.CENTER)
    if subtitle:
        add_textbox(slide4, x, y + Inches(1.05), Inches(3.8), Inches(0.3),
                    subtitle, font_size=14, color=GRAY, alignment=PP_ALIGN.CENTER)
    add_textbox(slide4, x + Inches(0.3), y + Inches(1.35), Inches(3.2), Inches(1.2),
                desc, font_size=13, color=GRAY, alignment=PP_ALIGN.CENTER)

add_slide_number(slide4, 4)

# ============================================================
# SLIDE 5: 使い方（4ステップ）
# ============================================================
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide5, WHITE)
add_header_bar(slide5, "基本的な使い方（4ステップ）", TEAL)

steps = [
    ("1", "ノートブック作成", "notebooklm.google.com に\nアクセスし、新しい\nノートブックを作成"),
    ("2", "ソースを追加", "PDF、URL、動画、音声\nなど最大300のソースを\nアップロード"),
    ("3", "AIに質問・指示", "チャットで質問したり、\n要約・分析・比較\nなどを依頼"),
    ("4", "成果物を活用", "音声・動画・スライド・\nマインドマップなど多様な\n形式で出力・共有"),
]

for i, (num, title, desc) in enumerate(steps):
    x = Inches(0.5 + i * 3.2)
    # Circle
    circle = add_circle(slide5, x + Inches(1.1), Inches(1.6), Inches(0.8), TEAL)
    tf = circle.text_frame
    p = tf.paragraphs[0]
    p.text = num
    p.font.size = Pt(28)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    add_textbox(slide5, x, Inches(2.6), Inches(3.0), Inches(0.5),
                title, font_size=20, color=DARK, bold=True, alignment=PP_ALIGN.CENTER)
    add_textbox(slide5, x, Inches(3.2), Inches(3.0), Inches(1.5),
                desc, font_size=14, color=GRAY, alignment=PP_ALIGN.CENTER)

    # Arrow
    if i < 3:
        add_textbox(slide5, x + Inches(2.8), Inches(1.75), Inches(0.6), Inches(0.5),
                    "→", font_size=30, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

# Tip box
tip_box = add_rounded_rect(slide5, Inches(0.6), Inches(5.5), Inches(12.1), Inches(1.2),
                           RGBColor(0xE0, 0xF2, 0xF1))
add_textbox(slide5, Inches(0.9), Inches(5.65), Inches(11.5), Inches(0.9),
            "💡 ポイント：NotebookLMはアップロードされた資料のみを情報源とするため、"
            "ハルシネーション（AIの誤情報生成）のリスクが大幅に低減されます。",
            font_size=15, color=RGBColor(0x00, 0x69, 0x5C))

add_slide_number(slide5, 5)

# ============================================================
# SLIDE 6: ビジネス活用事例
# ============================================================
slide6 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide6, WHITE)
add_header_bar(slide6, "ビジネス活用事例", PURPLE)

use_cases = [
    ("📋", "会議議事録の要約", "長時間の会議録から決定事項・担当者・\n期限・未決タスクを自動整理。\n議事録作成時間を大幅削減。"),
    ("🔎", "リサーチ・資料の要点抽出", "複数の資料やWebサイトを一括投入し、\n要点を自動抽出。\n調査レポート作成を効率化。"),
    ("📚", "社内ナレッジベースの構築", "部門別のAIアシスタントを手軽に作成。\n新人研修やFAQ対応にも活用可能。\nURL共有で簡単展開。"),
    ("📈", "企業調査・市場分析", "IR資料やWebサイトをアップロードして\n企業概要・事業内容を迅速分析。\n会議前の準備に最適。"),
]

for i, (icon, title, desc) in enumerate(use_cases):
    col = i % 2
    row = i // 2
    x = Inches(0.5 + col * 6.4)
    y = Inches(1.4 + row * 2.9)

    card = add_rounded_rect(slide6, x, y, Inches(6.0), Inches(2.5), BG_LIGHT)
    add_textbox(slide6, x + Inches(0.3), y + Inches(0.3), Inches(0.6), Inches(0.6),
                icon, font_size=36)
    add_textbox(slide6, x + Inches(1.1), y + Inches(0.35), Inches(4.5), Inches(0.4),
                title, font_size=20, color=DARK, bold=True)
    add_textbox(slide6, x + Inches(1.1), y + Inches(0.9), Inches(4.5), Inches(1.4),
                desc, font_size=14, color=GRAY)

add_slide_number(slide6, 6)

# ============================================================
# SLIDE 7: 強み
# ============================================================
slide7 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide7, WHITE)
add_header_bar(slide7, "NotebookLM の強み", INDIGO)

strengths = [
    ("🛡️", "ハルシネーションが少ない",
     "資料にない情報は「答えない」\n制限があるため、AIが勝手に\n情報を作り出すリスクが低い。\nビジネス利用に安心。",
     RGBColor(0xE8, 0xF0, 0xFE)),
    ("📎", "出典が明確",
     "回答には必ず元データの引用\nリンクが付くため、情報の根拠\nを即座に確認可能。レポートや\n提案書の信頼性向上に貢献。",
     RGBColor(0xE6, 0xF4, 0xEA)),
    ("🔒", "プライバシー保護",
     "アップロードデータはAIモデル\nのトレーニングに使用されない\n（Google公式発表）。\n企業の機密情報も安心。",
     RGBColor(0xFE, 0xF7, 0xE0)),
]

for i, (icon, title, desc, bg) in enumerate(strengths):
    x = Inches(0.5 + i * 4.2)
    card = add_rounded_rect(slide7, x, Inches(1.6), Inches(3.8), Inches(5.0), bg)
    add_textbox(slide7, x, Inches(2.0), Inches(3.8), Inches(0.7),
                icon, font_size=48, alignment=PP_ALIGN.CENTER)
    add_textbox(slide7, x, Inches(2.8), Inches(3.8), Inches(0.5),
                title, font_size=20, color=DARK, bold=True, alignment=PP_ALIGN.CENTER)
    add_textbox(slide7, x + Inches(0.3), Inches(3.5), Inches(3.2), Inches(2.5),
                desc, font_size=15, color=GRAY, alignment=PP_ALIGN.CENTER)

add_slide_number(slide7, 7)

# ============================================================
# SLIDE 8: 料金プラン
# ============================================================
slide8 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide8, WHITE)
add_header_bar(slide8, "料金プラン", YELLOW)

plans = [
    ("無料プラン", "¥0 / 月", BLUE,
     "✓ Googleアカウントで即利用開始\n✓ 基本的な要約・質問応答\n✓ 音声解説・動画解説\n✓ マインドマップ生成\n✓ 個人利用に最適"),
    ("NotebookLM Plus", "$19.99 / 月", GREEN,
     "✓ ソース量が無料版の6倍\n✓ ノートブック数が5倍\n✓ 年払い $200/年（約16%割引）\n✓ 優先サポート\n✓ チーム利用に推奨"),
    ("Enterprise", "要問合せ", RGBColor(0xE6, 0x51, 0x00),
     "✓ VPC Service Controls対応\n✓ IAMによるアクセス制御\n✓ データのリージョン限定保存\n✓ 管理者ダッシュボード\n✓ 大規模組織向け"),
]

for i, (name, price, color, features_text) in enumerate(plans):
    x = Inches(0.5 + i * 4.2)
    card = add_rounded_rect(slide8, x, Inches(1.5), Inches(3.8), Inches(5.2), WHITE)
    card.line.color.rgb = color
    card.line.width = Pt(2)

    if i == 1:  # Highlight Plus plan
        highlight_bg = add_rounded_rect(slide8, x, Inches(1.5), Inches(3.8), Inches(5.2),
                                        RGBColor(0xE6, 0xF4, 0xEA))
        highlight_bg.line.color.rgb = GREEN
        highlight_bg.line.width = Pt(2)

    add_textbox(slide8, x, Inches(1.8), Inches(3.8), Inches(0.5),
                name, font_size=22, color=color, bold=True, alignment=PP_ALIGN.CENTER)
    add_textbox(slide8, x, Inches(2.5), Inches(3.8), Inches(0.6),
                price, font_size=32, color=DARK, bold=True, alignment=PP_ALIGN.CENTER)
    add_textbox(slide8, x + Inches(0.4), Inches(3.4), Inches(3.0), Inches(3.0),
                features_text, font_size=14, color=GRAY)

add_slide_number(slide8, 8)

# ============================================================
# SLIDE 9: 注意点
# ============================================================
slide9 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide9, WHITE)
add_header_bar(slide9, "利用時の注意点", RED)

cautions = [
    ("⚠️", "機密情報の取り扱い",
     "顧客情報や契約書、個人情報など機密性の高い資料のアップロードは社内規定を確認の上、"
     "慎重に判断してください。情報管理ルールの明確化が必要です。"),
    ("🔄", "最新情報の反映にはソース更新が必要",
     "NotebookLMはアップロードされた資料のみを参照します。最新情報を反映するには、"
     "定期的にソースを更新・追加する運用が必要です。"),
    ("✅", "出力結果の確認は必須",
     "ハルシネーションは少ないものの、AIの出力は必ず人間が確認・検証してください。"
     "特に重要な意思決定に使用する場合は、出典元の原文も確認しましょう。"),
]

for i, (icon, title, desc) in enumerate(cautions):
    y = Inches(1.5 + i * 1.9)
    card = add_rounded_rect(slide9, Inches(0.6), y, Inches(12.1), Inches(1.6),
                            RGBColor(0xFF, 0xF3, 0xE0))
    # Yellow left border
    add_rect(slide9, Inches(0.6), y, Inches(0.08), Inches(1.6), YELLOW)

    add_textbox(slide9, Inches(0.9), y + Inches(0.2), Inches(0.6), Inches(0.5),
                icon, font_size=28)
    add_textbox(slide9, Inches(1.6), y + Inches(0.2), Inches(10.5), Inches(0.4),
                title, font_size=19, color=DARK, bold=True)
    add_textbox(slide9, Inches(1.6), y + Inches(0.75), Inches(10.5), Inches(0.7),
                desc, font_size=14, color=GRAY)

add_slide_number(slide9, 9)

# ============================================================
# SLIDE 10: まとめ & ネクストステップ
# ============================================================
slide10 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide10, DARK)

add_textbox(slide10, Inches(0.5), Inches(0.8), Inches(12.3), Inches(0.8),
            "まとめ & ネクストステップ", font_size=38, color=WHITE, bold=True,
            alignment=PP_ALIGN.CENTER)

next_steps = [
    ("🚀", "まず使ってみる", "無料でGoogleアカウント\nから即開始"),
    ("📄", "業務資料を投入", "議事録・マニュアル・\nレポートから試す"),
    ("👥", "チームで共有", "ノートブックをURL共有\nして活用拡大"),
]

for i, (icon, title, desc) in enumerate(next_steps):
    x = Inches(1.5 + i * 3.8)
    card = add_rounded_rect(slide10, x, Inches(2.2), Inches(3.2), Inches(2.8),
                            RGBColor(0x37, 0x47, 0x4F))
    add_textbox(slide10, x, Inches(2.5), Inches(3.2), Inches(0.5),
                icon, font_size=36, alignment=PP_ALIGN.CENTER)
    add_textbox(slide10, x, Inches(3.1), Inches(3.2), Inches(0.4),
                title, font_size=20, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    add_textbox(slide10, x, Inches(3.6), Inches(3.2), Inches(1.0),
                desc, font_size=14, color=RGBColor(0xBB, 0xBB, 0xBB),
                alignment=PP_ALIGN.CENTER)

add_textbox(slide10, Inches(0.5), Inches(5.5), Inches(12.3), Inches(0.6),
            "今日からNotebookLMで業務を変革しましょう！",
            font_size=24, color=YELLOW, bold=True, alignment=PP_ALIGN.CENTER)

add_textbox(slide10, Inches(0.5), Inches(6.8), Inches(12.3), Inches(0.4),
            "Google NotebookLM 社内研修資料  ｜  2026年3月",
            font_size=12, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

add_slide_number(slide10, 10)

# ============================================================
# Save
# ============================================================
output_path = "notebooklm-training-slides.pptx"
prs.save(output_path)
print(f"✅ PPTX saved: {output_path}")
